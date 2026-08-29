#!/usr/bin/env python3
"""
Build the Influence Hub lead-base template (xlsx).

Output: gtm/inflnc_leads_template.xlsx

Sheets:
  1. Leads        — the working list (scoring formula + statuses, 1 example row)
  2. ICP & Scoring — the qualification checklist and how the score is computed
  3. Lists (ref)  — dropdown source values (types, districts, statuses, channels)

Run:  python3 gtm/build_leads_template.py
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# Palette / styles
# ----------------------------------------------------------------------------
FONT = "Arial"
INK = "1A1A1A"
HEADER_FILL = PatternFill("solid", fgColor="16324F")   # deep blue
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
INPUT_FONT = Font(name=FONT, color="0000FF", size=10)   # blue = you fill this in
CALC_FONT = Font(name=FONT, color="1A1A1A", size=10)     # black = formula
TITLE_FONT = Font(name=FONT, bold=True, size=16, color=INK)
SUB_FONT = Font(name=FONT, italic=True, size=10, color="6B6B6B")
LEGEND_FILL = PatternFill("solid", fgColor="FFF7DB")     # soft yellow legend
EXAMPLE_FILL = PatternFill("solid", fgColor="EEF4FA")    # light blue example row
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()

# ============================================================================
# Sheet 3 first (reference lists) so validations can point at it
# ============================================================================
ref = wb.active
ref.title = "Lists (ref)"

LIST_TYPES = [
    "Specialty coffee", "Brunch / cafe", "Dessert / bakery", "Restaurant (dine-in)",
    "Healthy / bowls", "Bar / lounge", "Other",
]
LIST_DISTRICTS = [
    "JVC", "Business Bay", "Dubai Marina", "DIFC", "Downtown", "JLT", "Al Quoz",
    "City Walk", "Jumeirah", "Deira", "Abu Dhabi", "Sharjah", "Other",
]
LIST_SIZE = ["1 venue", "2-6 venues", "7+ venues"]
LIST_YESNO = ["Yes", "No", "Unknown"]
LIST_CHANNEL = ["WhatsApp", "Instagram DM", "Email", "Phone", "In person"]
LIST_ROLE = ["Owner", "Marketing / SMM", "GM / Manager", "Unknown"]
LIST_STATUS = [
    "0 - New", "1 - Enriched", "2 - Contacted", "3 - Replied", "4 - Demo booked",
    "5 - Trial started", "6 - Won", "X - Not a fit", "X - No response",
]

columns = {
    "A": ("Types", LIST_TYPES),
    "B": ("Districts", LIST_DISTRICTS),
    "C": ("Locations", LIST_SIZE),
    "D": ("Yes/No", LIST_YESNO),
    "E": ("Channel", LIST_CHANNEL),
    "F": ("Role", LIST_ROLE),
    "G": ("Status", LIST_STATUS),
}
for col, (title, values) in columns.items():
    c = ref[f"{col}1"]
    c.value = title
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center")
    for i, v in enumerate(values, start=2):
        cell = ref[f"{col}{i}"]
        cell.value = v
        cell.font = Font(name=FONT, size=10)
    ref.column_dimensions[col].width = 20

# named ranges helper (absolute refs into this sheet)
def ref_range(col, n):
    return f"'Lists (ref)'!${col}$2:${col}${n+1}"

RANGE_TYPES = ref_range("A", len(LIST_TYPES))
RANGE_DISTRICTS = ref_range("B", len(LIST_DISTRICTS))
RANGE_SIZE = ref_range("C", len(LIST_SIZE))
RANGE_YESNO = ref_range("D", len(LIST_YESNO))
RANGE_CHANNEL = ref_range("E", len(LIST_CHANNEL))
RANGE_ROLE = ref_range("F", len(LIST_ROLE))
RANGE_STATUS = ref_range("G", len(LIST_STATUS))

# ============================================================================
# Sheet 1: Leads
# ============================================================================
ws = wb.create_sheet("Leads", 0)

# Column model: (header, width, kind)  kind in {input, calc}
COLS = [
    ("Venue name",        24, "input"),
    ("Type",              18, "input"),
    ("District",          14, "input"),
    ("Locations",         12, "input"),
    ("IG handle",         18, "input"),
    ("IG followers",      12, "input"),
    ("Posts UGC?",        12, "input"),   # do THEY already get lots of creator content?
    ("Google rating",     12, "input"),
    ("Google reviews",    13, "input"),
    ("New (<6 mo)?",      12, "input"),
    ("Website",           22, "input"),
    ("Phone",             16, "input"),
    ("WhatsApp",          16, "input"),
    ("Email",             22, "input"),
    ("Decision maker",    18, "input"),
    ("Role",              16, "input"),
    ("Best channel",      14, "input"),
    ("Score",             8,  "calc"),
    ("Tier",              8,  "calc"),
    ("Status",            16, "input"),
    ("Last touch",        13, "input"),
    ("Next step",         22, "input"),
    ("Notes",             34, "input"),
]

TITLE_ROW = 1
LEGEND_ROW = 2
HEADER_ROW = 3
FIRST_DATA_ROW = 4
LAST_DATA_ROW = 503   # 500 lead rows

n_cols = len(COLS)
last_col_letter = get_column_letter(n_cols)

# Title + legend
ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=n_cols)
t = ws.cell(TITLE_ROW, 1, "Influence Hub — Lead base (F&B venues)")
t.font = TITLE_FONT

ws.merge_cells(start_row=LEGEND_ROW, start_column=1, end_row=LEGEND_ROW, end_column=n_cols)
legend = ws.cell(
    LEGEND_ROW, 1,
    "Fill BLUE columns. Score & Tier are auto (black). Work top-down by Tier A → B → C. "
    "Row 4 is an example — overwrite or delete it.",
)
legend.font = SUB_FONT
legend.fill = LEGEND_FILL
legend.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[LEGEND_ROW].height = 28

# Header row
for i, (name, width, kind) in enumerate(COLS, start=1):
    c = ws.cell(HEADER_ROW, i, name)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[HEADER_ROW].height = 30

# Column-letter lookup by header
IDX = {name: get_column_letter(i) for i, (name, _, _) in enumerate(COLS, start=1)}

def score_formula(r):
    """
    A/B/C scoring. Higher = hotter. Points:
      +3  active IG (>=3000 followers) but does NOT already get lots of UGC  -> core fit
      +2  new venue (<6 months)
      +1  small operator (1-6 venues)
      +1  visual segment (coffee / brunch / dessert)
      +1  strong social proof (rating >= 4.3 AND reviews >= 50)
      +1  we have a direct decision-maker contact (WhatsApp or Email + name)
    """
    ig = IDX["IG followers"]
    ugc = IDX["Posts UGC?"]
    new = IDX["New (<6 mo)?"]
    loc = IDX["Locations"]
    typ = IDX["Type"]
    rat = IDX["Google rating"]
    rev = IDX["Google reviews"]
    wa = IDX["WhatsApp"]
    em = IDX["Email"]
    dm = IDX["Decision maker"]
    name = IDX["Venue name"]
    return (
        f'=IF({name}{r}="","",'
        f'IF(AND(N({ig}{r})>=3000,{ugc}{r}<>"Yes"),3,0)'
        f'+IF({new}{r}="Yes",2,0)'
        f'+IF(OR({loc}{r}="1 venue",{loc}{r}="2-6 venues"),1,0)'
        f'+IF(OR({typ}{r}="Specialty coffee",{typ}{r}="Brunch / cafe",{typ}{r}="Dessert / bakery"),1,0)'
        f'+IF(AND(N({rat}{r})>=4.3,N({rev}{r})>=50),1,0)'
        f'+IF(AND(OR({wa}{r}<>"",{em}{r}<>""),{dm}{r}<>""),1,0)'
        ")"
    )

def tier_formula(r):
    s = IDX["Score"]
    # A >=6, B 4-5, C 1-3, blank if no score yet
    return f'=IF({s}{r}="","",IF({s}{r}>=6,"A",IF({s}{r}>=4,"B","C")))'

# Example row (row 4)
example = {
    "Venue name": "Marlow Coffee (example)",
    "Type": "Specialty coffee",
    "District": "JVC",
    "Locations": "1 venue",
    "IG handle": "@marlowcoffee",
    "IG followers": 8200,
    "Posts UGC?": "No",
    "Google rating": 4.6,
    "Google reviews": 180,
    "New (<6 mo)?": "Yes",
    "Website": "marlowcoffee.ae",
    "Phone": "+971 50 000 0000",
    "WhatsApp": "+971 50 000 0000",
    "Email": "hi@marlowcoffee.ae",
    "Decision maker": "Sara (owner)",
    "Role": "Owner",
    "Best channel": "WhatsApp",
    "Status": "1 - Enriched",
    "Last touch": "2026-08-25",
    "Next step": "Send opener + 2 profiles",
    "Notes": "Nice interior, weak UGC flow — perfect trial-month pitch",
}

for name, _, kind in COLS:
    col = IDX[name]
    cell = ws[f"{col}{FIRST_DATA_ROW}"]
    if kind == "calc":
        cell.value = score_formula(FIRST_DATA_ROW) if name == "Score" else tier_formula(FIRST_DATA_ROW)
        cell.font = CALC_FONT
    else:
        cell.value = example.get(name, "")
        cell.font = INPUT_FONT
    cell.fill = EXAMPLE_FILL
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center", wrap_text=(name == "Notes"))

# Remaining empty rows: only the calc formulas prefilled
for r in range(FIRST_DATA_ROW + 1, LAST_DATA_ROW + 1):
    for name, _, kind in COLS:
        col = IDX[name]
        cell = ws[f"{col}{r}"]
        cell.border = BORDER
        if kind == "calc":
            cell.value = score_formula(r) if name == "Score" else tier_formula(r)
            cell.font = CALC_FONT
        else:
            cell.font = INPUT_FONT

# Number / date formats
for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
    ws[f'{IDX["IG followers"]}{r}'].number_format = "#,##0"
    ws[f'{IDX["Google reviews"]}{r}'].number_format = "#,##0"
    ws[f'{IDX["Google rating"]}{r}'].number_format = "0.0"
    ws[f'{IDX["Last touch"]}{r}'].number_format = "yyyy-mm-dd"

# Freeze header + first columns
ws.freeze_panes = f"C{FIRST_DATA_ROW}"

# Dropdown validations
def add_dv(col_name, source):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    ws.add_data_validation(dv)
    col = IDX[col_name]
    dv.add(f"{col}{FIRST_DATA_ROW}:{col}{LAST_DATA_ROW}")

add_dv("Type", RANGE_TYPES)
add_dv("District", RANGE_DISTRICTS)
add_dv("Locations", RANGE_SIZE)
add_dv("Posts UGC?", RANGE_YESNO)
add_dv("New (<6 mo)?", RANGE_YESNO)
add_dv("Role", RANGE_ROLE)
add_dv("Best channel", RANGE_CHANNEL)
add_dv("Status", RANGE_STATUS)

ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col_letter}{LAST_DATA_ROW}"

# ============================================================================
# Sheet 2: ICP & Scoring
# ============================================================================
doc = wb.create_sheet("ICP & Scoring", 1)
doc.column_dimensions["A"].width = 4
doc.column_dimensions["B"].width = 52
doc.column_dimensions["C"].width = 10
doc.column_dimensions["D"].width = 60

def h(row, text):
    doc.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    c = doc.cell(row, 2, text)
    c.font = Font(name=FONT, bold=True, size=13, color="16324F")

def line(row, b, c="", d=""):
    doc.cell(row, 2, b).font = Font(name=FONT, size=10, color=INK)
    doc.cell(row, 3, c).font = Font(name=FONT, size=10, bold=True, color=INK)
    dd = doc.cell(row, 4, d)
    dd.font = Font(name=FONT, size=10, color="4A4A4A")
    dd.alignment = Alignment(wrap_text=True, vertical="top")

r = 1
doc.cell(r, 2, "ICP & Scoring — how to qualify a lead").font = TITLE_FONT
r += 2
h(r, "Ideal Customer Profile (the 5-point checklist)"); r += 1
for b, d in [
    ("Type: coffee / brunch / dessert / visual dine-in", "Places where visuals & footfall matter. Skip fast food, dark kitchens, delivery-only."),
    ("Stage: newly opened (<6 mo) OR active IG but irregular UGC", "They need content + guests now — the sharpest pain your offer solves."),
    ("Size: 1-6 venues", "Owner/marketing decides fast, no tender. Matches the form's segment."),
    ("Geo: 2-3 Dubai districts first (JVC, Business Bay, Marina, DIFC)", "Density = reusable cases and warm referrals before you scale."),
    ("Reachable decision maker", "You can find owner / SMM by name + a WhatsApp or email — not a generic info@ line."),
]:
    line(r, "•  " + b, "", d); r += 1
r += 1

h(r, "Scoring (auto-computed in the Leads sheet)"); r += 1
line(r, "Signal", "Points", "Why it matters");
for cc in (2,3,4):
    doc.cell(r, cc).font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
    doc.cell(r, cc).fill = HEADER_FILL
r += 1
for b, pts, d in [
    ("Active IG (≥3000 followers) AND does NOT already get lots of UGC", "+3", "Core fit: audience exists, but the creator flow is missing — exactly your product."),
    ("New venue (<6 months)", "+2", "Highest urgency — they're building buzz right now."),
    ("Small operator (1-6 venues)", "+1", "Fast, single decision maker."),
    ("Visual segment (coffee / brunch / dessert)", "+1", "Best content ROI for creators."),
    ("Strong social proof (rating ≥4.3 AND ≥50 reviews)", "+1", "Real, working venue — worth a creator's visit."),
    ("Direct decision-maker contact (name + WhatsApp/email)", "+1", "You can actually start a 1:1 conversation."),
]:
    line(r, b, pts, d); r += 1
r += 1
line(r, "Tier A", "≥ 6", "Work first. Personal 1:1 WhatsApp/IG opener + 2 real creator profiles.")
r += 1
line(r, "Tier B", "4 - 5", "Second wave. Personalised but lighter touch.")
r += 1
line(r, "Tier C", "1 - 3", "Email drip / nurture only. Don't spend manual time yet.")
r += 2

h(r, "Compliance (UAE)"); r += 1
for b in [
    "UAE PDPL + anti-spam rules apply. Collect only public business contacts.",
    "WhatsApp is the main channel — but first touch is 1:1 and personal, never a bulk broadcast.",
    "Email: keep an opt-out. No cold bulk WhatsApp broadcasts to an unconsented list.",
    "Quality > volume: 50 enriched Tier-A leads beat 1000 raw 'name + phone' rows.",
]:
    line(r, "•  " + b); r += 1

for sheet in (doc, ref):
    sheet.sheet_view.showGridLines = False

OUT = "gtm/inflnc_leads_template.xlsx"
wb.save(OUT)
print("saved", OUT)
